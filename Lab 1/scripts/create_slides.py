#!/usr/bin/env python3
"""Build the editable lab presentation from the measured project artifacts.

Requires python-pptx and Pillow. Run from any directory.
"""
from pathlib import Path
import csv
import json
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'slides'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY='12283E'; BLUE='176B9B'; ORANGE='D96B27'; GREEN='23826B'; MUTED='566779'; LIGHT='EEF3F7'; WHITE='FFFFFF'
selection=json.loads((ROOT/'results/full/selection.json').read_text())
S=selection['selected_s']
stats={r['algorithm']:r for r in csv.DictReader((ROOT/'results/full/validation_summary.csv').open())}
findings=(ROOT/'results/full/findings.md').read_text()
speed=re.search(r'\*\*(\d+\.\d+)×\*\*',findings).group(1)
change=100*(float(stats['hybrid']['comparisons_median'])/float(stats['merge']['comparisons_median'])-1)
blueprint=(OUT/'blueprint.md').read_text()

def box(slide,x,y,w,h,fill,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(fill)
    shape.line.fill.background() if not line else None
    if line: shape.line.color.rgb=RGBColor.from_string(line)
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    return shape

def text(slide,words,x,y,w,h,size=22,color=NAVY,bold=False,font='Arial',align=None):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.01); tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(words.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(8)
        if align is not None: p.alignment=align
    return sh

def slide(title,kicker='SC2001 / PROJECT 1',dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else WHITE)
    fg=WHITE if dark else NAVY
    text(s,kicker,.55,.25,12,.3,11,ORANGE,True)
    text(s,title,.55,.75,12.25,.85,30,fg,True)
    box(s,.55,7.08,12.2,.012,MUTED if dark else 'DDE4EA')
    text(s,'MERGE SORT + INSERTION SORT',.55,7.18,8,.2,9,'B2C3D2' if dark else MUTED)
    text(s,f'{len(prs.slides):02d}',12.1,7.14,.6,.25,11,fg,align=PP_ALIGN.RIGHT)
    return s

def note(s,body,main=None):
    if main:
        match=re.search(rf'## Slide {main} —.*?(?=\n## |\Z)',blueprint,re.S)
        if match: body+='\n\n'+match.group(0)
    s.notes_slide.notes_text_frame.text=body

def picture(s,name,x,y,w,h):
    path=ROOT/'figures'/name
    iw,ih=Image.open(path).size; scale=min(w/iw,h/ih)
    width,height=iw*scale,ih*scale
    s.shapes.add_picture(str(path), Inches(x+(w-width)/2), Inches(y+(h-height)/2),width=Inches(width),height=Inches(height))

def callout(s,title,body,x,y,w,h=1.15,color=BLUE):
    box(s,x,y,w,h,LIGHT)
    box(s,x,y,.055,h,color)
    text(s,title,x+.18,y+.12,w-.36,.4,20,color,True)
    text(s,body,x+.18,y+.55,w-.36,h-.6,16,MUTED)

# 1 — research question, framed by actual findings.
s=slide('Can a small-array switch make merge sort faster?','SC2001 / PROJECT 1 / 0:00–0:30',True)
text(s,'Hybrid merge sort\n+ insertion sort',.65,1.9,8,1.55,40,WHITE,True)
text(s,'Choose a threshold S. Measure both CPU time and key comparisons.',.65,3.8,7.2,1,24,'C8D5DF')
box(s,9.1,2.0,3.5,3.15,BLUE)
text(s,f'{speed}×',9.3,2.4,3.1,.9,52,WHITE,True,align=PP_ALIGN.CENTER)
text(s,'measured speedup\non fresh 10m-element arrays',9.35,3.5,3,.95,20,WHITE,align=PP_ALIGN.CENTER)
text(s,'n = input length     S = insertion-sort threshold',.65,5.6,11,.45,21,WHITE)
text(s,'Research question: does the fastest S also minimize comparisons?',.65,6.25,12,.4,22,'C8D5DF')
note(s,'30 seconds. Explain the research question. The headline is a measured result on this implementation, not a universal claim. Add team names and lab group if desired before presenting.',1)

# 2 — readable pseudocode and editable tree.
s=slide('Sort small leaves; merge the sorted halves','ALGORITHM & CORRECTNESS / (a) / 0:30–1:40')
box(s,.55,1.8,6.4,4.75,LIGHT)
code='hybrid(A, left, right, S):\n  if right-left <= 1: return\n  if right-left <= S:\n    insertionSort(A, left, right)\n    return\n  mid = left + (right-left)/2\n  hybrid(A, left, mid, S)\n  hybrid(A, mid, right, S)\n  merge(A, left, mid, right)'
text(s,code,.8,2.03,6,4.3,18,NAVY,font='Menlo')
nodes=[(9.85,2.5,'16',BLUE),(8.55,3.45,'8',BLUE),(11.15,3.45,'8',BLUE),
       (7.85,4.4,'4',ORANGE),(9.25,4.4,'4',ORANGE),(10.65,4.4,'4',ORANGE),(12.05,4.4,'4',ORANGE)]
for parent,child in [(0,1),(0,2),(1,3),(1,4),(2,5),(2,6)]:
    x1,y1,*_=nodes[parent]; x2,y2,*_=nodes[child]
    line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1+.3),Inches(x2),Inches(y2-.3)); line.line.color.rgb=RGBColor.from_string(MUTED)
for x,y,label,color in nodes:
    box(s,x-.4,y-.28,.8,.58,color); text(s,label,x-.35,y-.22,.7,.4,20,WHITE,True,align=PP_ALIGN.CENTER)
text(s,'n=16, S=4',7.5,1.7,5,.4,21,NAVY,True,align=PP_ALIGN.CENTER)
text(s,'Insertion sort at the four leaves',7.35,4.95,5.35,.45,20,ORANGE,True,align=PP_ALIGN.CENTER)
text(s,'Sorted leaves → sorted halves → sorted whole\nStable ties • one reusable buffer',7.35,5.45,5.35,1,18,MUTED,align=PP_ALIGN.CENTER)
note(s,'70 seconds. Tree nodes show subarray lengths. Correctness follows from insertion-sort invariant and induction through stable merge. Half-open intervals exclude the right endpoint.',2)

# 3 — mathematical argument, no unexplained exact predictive fit.
s=slide('The threshold trades leaf work for merge levels','THEORETICAL ANALYSIS / (c) / 1:40–2:40')
callout(s,'Insertion-sort leaves','About n/S leaves × O(S²) work per leaf',.65,1.85,5.9,1.35,ORANGE)
callout(s,'Merge levels','About log₂(n/S) levels × O(n) work per level',6.8,1.85,5.9,1.35,BLUE)
text(s,'T(n,S) = O(nS + n log₂(n/S))',.8,3.65,11.8,.7,36,NAVY,True,align=PP_ALIGN.CENTER)
text(s,'Worst-case bound for 1 ≤ S ≤ n; not an exact random-input count.',1,4.5,11.3,.5,20,MUTED,align=PP_ALIGN.CENTER)
for x,title,body in [(.65,'Fixed S','Θ(n log n) worst case'),(4.8,'S = 1','Ordinary merge sort'),(8.95,'S ≥ n','Insertion sort: O(n²) worst case')]: callout(s,title,body,x,5.45,3.75,1.15)
note(s,'60 seconds. Auxiliary storage O(n), plus O(log(n/S)+1) recursive stack depth for 1≤S≤n. Do not claim differentiation of the bound identifies a universal threshold.',3)

# 4 — experiment flow: make the data journey explicit.
s=slide('Experiment flow: from generated arrays to a fair comparison','PART (b) → (c) → (d) / 2:40–3:30')
text(s,'Part (b) creates the inputs. We then use those inputs to choose S and test it on fresh data.',.7,1.52,12,.38,18,MUTED)

# Three stages with explicit “what” and “why”, connected by directional arrows.
stages=[
    ('01','GENERATE','Uniform integers in [1, 10⁷]\n9 sizes: 1,000 → 10,000,000','Why: isolate the effect of n',BLUE),
    ('02','TUNE S','Sweep candidate thresholds\n5 seeds × 3 timing repeats\nChoose the lowest median CPU time','Why: find a fast S before testing it',ORANGE),
    ('03','VALIDATE','Lock S = 64 before testing\n5 fresh 10m-element arrays\nCompare merge sort vs hybrid','Why: avoid choosing S on test data',GREEN),
]
for i,(num,title,body,why,color) in enumerate(stages):
    x=.65+i*4.22
    box(s,x,2.08,3.72,2.55,LIGHT,line='D5E0E8',radius=True)
    box(s,x+.2,2.3,.55,.55,color)
    text(s,num,x+.2,2.38,.55,.25,15,WHITE,True,align=PP_ALIGN.CENTER)
    text(s,title,x+.92,2.34,2.5,.35,19,color,True)
    text(s,body,x+.22,2.98,3.28,1.18,15,NAVY)
    text(s,why,x+.22,4.24,3.28,.25,13,MUTED,True)
    if i<2:
        line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+3.76),Inches(3.35),Inches(x+4.12),Inches(3.35))
        line.line.color.rgb=RGBColor.from_string(MUTED); line.line.width=Pt(2.5)
        line.line.end_arrowhead=True

# Controls are subordinate to the flow, but remain visible for marks on methodology.
box(s,.65,5.05,12.0,1.18,'F7FAFC',line='D5E0E8',radius=True)
text(s,'CONTROLS HELD CONSTANT',.9,5.23,3.1,.28,14,BLUE,True)
text(s,'Same array for every candidate',.9,5.67,3.55,.3,16,NAVY,True)
text(s,'Warm-up + shuffled run order',4.67,5.67,3.35,.3,16,NAVY,True)
text(s,'Sort only is timed; counters run separately',8.25,5.67,4.0,.3,16,NAVY,True)
text(s,'Every output checked against std::sort  •  Apple M5 / C++17 / -O3',.75,6.55,12,.3,15,MUTED)
note(s,'50 seconds. Walk left to right: first explain what Part (b) generates, then explain that threshold tuning is a separate stage, then explain why validation uses fresh arrays after S is locked. The lower strip gives the fairness controls without interrupting the main story. CPU time is process CPU via std::clock; generation, copying, allocation, checks and CSV output are excluded. Error bars later are IQRs, not confidence intervals.',4)

# 5 — fixed S graph.
s=slide('Fixed S: comparison growth agrees with n log n','EMPIRICAL ANALYSIS / (c)(i) / 3:30–4:25')
picture(s,'01_fixed_s.png',.5,1.65,12.3,4.85)
text(s,'S=32  •  normalized ratio 1.042–1.343  •  finite-size leaf effects explain the variation',.7,6.57,12,.35,17,MUTED)
note(s,'55 seconds. The right panel is not perfectly flat: actual leaf sizes vary as n changes, changing insertion work per element. The dashed reference is scaled to the largest point and is not an independently predicted count. Data is consistent with the bound; the derivation establishes it.',5)

# 6 — two required views, comparison plot and focused CPU view.
s=slide('Fastest runtime does not mean fewest comparisons','THRESHOLD SELECTION / (c)(ii)–(iii) / 4:25–5:45')
picture(s,'02_fixed_n.png',.45,1.65,6.25,3.8)
picture(s,'deck_cpu_10m.png',6.7,1.65,6.2,3.8)
callout(s,f'Selected S={S}','Minimum comparisons at 10m: S=1',.65,5.55,3.8,1.1,ORANGE)
callout(s,'Coarse winners by n','10k: 48  ·  100k: 96  ·  1m: 32  ·  10m: 64',4.65,5.55,7.95,1.1)
note(s,'80 seconds. Left: fixed n=100k key comparisons. Right: coarse CPU timings at 10m. All four CPU curves and the refinement graph are in backup slides. The selected S=64 remains the best measured candidate after refining 48–96. S=39–75 generates the same leaves at 10m, so exact ranking within that plateau is not algorithmically meaningful. Later refinement timings differ from reused coarse timings; do not hide this phase limitation.',6)

# 7 — held-out outcome.
s=slide(f'Hybrid S={S}: {speed}× speedup with {change:.1f}% more comparisons','HELD-OUT COMPARISON / (d) / 5:45–6:45')
picture(s,'06_validation.png',.55,1.6,12.2,5.0)
text(s,'Speedup = median of per-seed time ratios  •  IQR: 1.138–1.152×  •  5 fresh datasets',.7,6.65,12,.3,16,MUTED)
note(s,'60 seconds. Median CPU: original 0.443770s; hybrid 0.386277s. Median comparisons: 220,101,526 vs 281,233,315. Speedup uses paired seed ratios, not the ratio of overall medians. The result supports a runtime benefit here despite greater comparison work; no universal claim.',7)

# 8 — live demo plus takeaways.
s=slide('Correctness, scaling, and a measured runtime benefit','DEMONSTRATION & CONCLUSION / 6:45–8:00',True)
box(s,.65,1.85,6.05,3.85,'203C54')
text(s,'LIVE DEMO',.9,2.1,5.5,.3,14,'79C4E6',True)
text(s,'./build/demo',.9,2.65,5.5,.45,25,WHITE,font='Menlo')
text(s,'Input:  7 2 5 1 6 3 4 2\nOutput: 1 2 2 3 4 5 6 7',.9,3.45,5.5,1,20,WHITE,font='Menlo')
text(s,'Merge / S=1: 17 comparisons\nS=4: 19  •  S=8: 22',.9,4.75,5.5,.8,20,'C8D5DF')
text(s,'01  Correctness checks passed\n02  Fixed-S growth fits the analysis\n03  Faster runtime, more comparisons',7.1,2.05,5.5,2.7,24,WHITE,True)
text(s,'Limit: one implementation and machine;\nthe exact best S is sensitive to timing variation.',7.1,5.05,5.4,1.15,20,'C8D5DF')
text(s,'Questions  /  2 minutes',.75,6.3,11.8,.5,25,'79C4E6',True)
note(s,'75 seconds: about 35s for demo and 40s for conclusions. Do not run the full benchmark live. Saved fallback: slides/demo-output.txt. End at eight minutes and invite Q&A.',8)

# Appendix: full curves, refinement, counting, checks, method, Q&A.
s=slide('CPU time across four input sizes','BACKUP A / THRESHOLD SENSITIVITY')
picture(s,'03_threshold_cpu.png',1.3,1.55,10.7,5.45)
note(s,'Coarse runtime winners differ across sizes. IQRs overlap for several candidates. Threshold selection is empirical and environment-specific.')
s=slide('Refinement reveals run-phase variation','BACKUP B / WHY THE EXACT OPTIMUM IS UNCERTAIN')
picture(s,'05_refinement.png',.6,1.6,12.1,4.95)
text(s,'At n=10m, S=39–75 gives identical leaves (38 or 39 elements). Tiny ranks are not robust.',.7,6.65,12,.3,16,MUTED)
note(s,'Orange markers reuse earlier coarse observations; blue points were measured in the later fine sweep. Their systematic timing difference within the same recursion partition is evidence of a measurement-phase limitation, not threshold-induced algorithmic work. The held-out S=64 vs original result is the basis for the runtime claim.')
s=slide('Key comparisons per element across input sizes','BACKUP C / COMPARISON COST')
picture(s,'04_threshold_comparisons.png',1,1.55,11.3,5.45)
note(s,'The curves cross because per-element insertion work depends on actual recursive leaf lengths, not just S. Counts are normalized by n here, not by n log n. Same-partition thresholds have identical comparisons on each seed.')
s=slide('Count data comparisons, including the final false one','BACKUP D / COUNTING DEFINITIONS')
callout(s,'Insertion example: [2, 1, 3]','2 > 1 → true; shift 2\n2 > 3 → false; stop\nTotal: 2 key comparisons',.65,1.9,5.85,2.3,ORANGE)
callout(s,'Merge example: [1, 3] + [2, 4]','1 ≤ 2; 3 ≤ 2; 3 ≤ 4\nCopy remaining 4 without comparing\nTotal: 3 key comparisons',6.75,1.9,5.85,2.3,BLUE)
text(s,'Excluded: index checks, loop bounds, copying, and exhaustion checks.',.8,4.8,11.8,.65,23,NAVY,True)
text(s,'Count=true records uint64 comparisons.\nCount=false removes counter increments at compile time.',.8,5.65,11.5,1,21,MUTED)
note(s,'Insertion uses j>left && previous>key. Short circuit at the left boundary does not evaluate another key comparison. Sorted or all-equal length four: 3 comparisons; reverse length four: 6.')
s=slide('Correctness and complete experiment coverage','BACKUP E / VERIFICATION')
callout(s,'2,020 recorded passes','505 configuration/seed groups; all required sizes and repeats',.65,1.85,12,1.25,GREEN)
text(s,'Algorithm checks',.8,3.55,5.4,.4,24,BLUE,True)
text(s,'Empty, singleton, odd lengths\nSorted, reversed, equal, duplicate keys\nRandom arrays and signed extremes\nS=1 exactly matches original counts',.8,4.15,5.65,2.25,21)
text(s,'Independent checks',7,3.55,5.3,.4,24,BLUE,True)
text(s,'Exact std::sort reference equality\nAddress + undefined-behavior sanitizers\nCSV coverage audit\nThreshold recomputed from saved data',7,4.15,5.55,2.25,21)
note(s,'Tests cover lengths 0–260 over random trials and multiple thresholds. Stable behavior follows from strict-greater insertion and left-first merge ties; integer reference checks alone do not test identity ordering among equal keys. Source: results/full/verification.md and audit.txt.')
s=slide('Reproducible measurements; explicit limitations','BACKUP F / METHOD DETAILS')
text(s,'Data & timing',.75,1.9,5.7,.45,24,BLUE,True)
text(s,'mt19937 + unbiased rejection sampling\nTuning: 101, 202, 303, 404, 505\nValidation: 1101, 1202, 1303, 1404, 1505\nThree timed repetitions per seed\nCPU: std::clock; elapsed: steady_clock',.75,2.55,6,2.8,19)
text(s,'Interpretation limits',7.05,1.9,5.5,.45,24,ORANGE,True)
text(s,'Five seeds; IQR is not a confidence interval\nCoarse and fine phases occur at different times\nSmall inputs are timer-resolution sensitive\nOne machine and uniform input distribution\nNo alternative-distribution benchmarks claimed',7.05,2.55,5.6,3.3,19)
text(s,'Regenerate: make test → scripts/run_experiments.py → scripts/plot_results.py',.75,6.15,12,.6,18,MUTED)
note(s,'Do not mix smoke-test outputs with full experiment evidence. The sorting auxiliary array is O(n); the benchmark keeps four integer arrays, about 160 MB at n=10m. Fresh seeds prevent retuning to held-out results.')
s=slide('Questions we should all be able to answer','BACKUP G / Q&A')
questions=[('Why can insertion sort help?','Quadratic work stays in small leaves; recursion and copying costs fall.'),('Why do thresholds form plateaus?','Several S values stop at the same recursive subarray sizes.'),('Why not minimize comparisons?','CPU time also includes movement and control overhead.'),('Is S=64 universally optimal?','No. It is the best observed candidate under this implementation and run.'),('Does the graph prove the complexity?','No. The derivation establishes the bound; measurements test consistency.')]
for i,(q,a) in enumerate(questions):
    y=1.75+i*.99
    text(s,q,.75,y,11.9,.35,21,BLUE,True)
    text(s,a,.75,y+.4,11.9,.4,18,MUTED)
note(s,'All team members should rehearse answering these questions. Sources: supplied Project 1.pdf and info.pdf; implementation src/sorting.hpp; experimental evidence results/full/measurements.csv. Eight main slides total eight minutes; seven backup slides are for Q&A only.')

prs.core_properties.title='SC2001 Project 1 — Hybrid Merge Sort'
prs.core_properties.subject='Algorithm correctness, complexity, threshold tuning, and held-out performance'
prs.core_properties.author='SC2001 Lab Team'
prs.core_properties.keywords='SC2001, merge sort, insertion sort, hybrid, experiments'
path=OUT/'SC2001_Project_1.pptx'; prs.save(path)
assert len(prs.slides)==15
for idx,sl in enumerate(prs.slides,1):
    assert sl.notes_slide.notes_text_frame.text.strip(), f'Missing notes: {idx}'
    for shape in sl.shapes:
        assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+10000 and shape.top+shape.height<=prs.slide_height+10000, f'Out of bounds: slide {idx}'
print(f'Created {path}: 8 main slides + 7 backups; notes on all slides.')
