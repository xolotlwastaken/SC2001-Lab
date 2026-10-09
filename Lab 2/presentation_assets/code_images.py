"""Render exact notebook excerpts as high-resolution, syntax-colored code images."""
import ast
import io
import keyword
import tokenize
from PIL import Image, ImageDraw, ImageFont
from pptx.util import Inches


def add_code_slides(prs, nb, assets, slide, text, band):
    functions = {}
    for cell_number, cell in enumerate(nb['cells'], 1):
        if cell['cell_type'] != 'code':
            continue
        source = ''.join(cell['source'])
        for node in ast.walk(ast.parse(source)):
            if isinstance(node, ast.FunctionDef):
                functions[node.name] = (cell_number, source.splitlines(), node.lineno, node.end_lineno)

    def excerpt(name, start=None, end=None):
        cell, lines, first, last = functions[name]
        first = first if start is None else start
        last = last if end is None else end
        return [(i, lines[i-1]) for i in range(first, last+1)], cell

    def render(rows, name):
        font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Courier New.ttf', 37)
        char = font.getlength('M')
        line_height = 53
        w = int(150 + max(len(line) for _, line in rows)*char)
        h = 78 + len(rows)*line_height
        im = Image.new('RGB', (w, h), '#12283E')
        draw = ImageDraw.Draw(im)
        for j,(number,line) in enumerate(rows):
            y=38+j*line_height
            draw.text((22,y),str(number) if number else '',font=font,fill='#71869B')
            draw.text((115,y),line,font=font,fill='#E8EFF5')
        # Tokenize the original text; tolerate excerpts that end mid-block.
        code='\n'.join(line for _,line in rows)+'\n'
        try:
            for token in tokenize.generate_tokens(io.StringIO(code).readline):
                kind,value,(r,c),(er,ec),_=token
                color = None
                if kind==tokenize.COMMENT: color='#A3B5C7'
                elif kind==tokenize.STRING: color='#B6DF9A'
                elif kind==tokenize.NUMBER: color='#F3B571'
                elif kind==tokenize.NAME and keyword.iskeyword(value): color='#75D4FF'
                if color and r==er:
                    draw.text((115+c*char,38+(r-1)*line_height),value,font=font,fill=color)
        except (tokenize.TokenError, IndentationError):
            pass
        path=assets/f'code_{name}.png'
        im.save(path)
        return path, w, h

    original=list(prs.slides._sldIdLst)
    additions={}
    specs=[
        (3,'Actual code: graph conversions','GRAPH REPRESENTATION',
         'conversions',[('to_adjacency_matrix',None,None),('to_adjacency_list',None,None)],
         'Matrix allocates V² cells; the list allocates V lists and appends E edges.'),
        (4,'Actual code: the shared relaxation function','ALGORITHM',
         'relax',[('relax',None,None)],
         'A successful relaxation updates both the distance and its parent.'),
        (7,'Actual code: Part A selects the next vertex','PART A • CODE 1 / 2',
         'matrix_queue',[('dijkstra_matrix',1,13)],
         'The unsorted queue scans its remaining vertices to find the minimum.'),
        (7,'Actual code: Part A scans the matrix row','PART A • CODE 2 / 2',
         'matrix_edges',[('dijkstra_matrix',14,23)],
         'Every extracted vertex scans all n columns, including absent edges.'),
        (12,'Actual code: Part B uses adjacency lists','PART B • CODE 1 / 3',
         'list',[('dijkstra_list',None,None)],
         'Only existing neighbors are visited; successful updates call decrease-key.'),
        (12,'Actual code: extracting the heap minimum','PART B • CODE 2 / 3',
         'extract',[('extract_min',None,None)],
         'Replace the root with the last entry, then sift down to restore heap order.'),
        (12,'Actual code: decreasing a heap priority','PART B • CODE 3 / 3',
         'decrease',[('decrease_key',None,None)],
         'position[vertex] finds the heap entry in O(1); sift-up takes O(log V).'),
    ]
    for after,title,section,name,parts,takeaway in specs:
        rows=[]; refs=[]
        for func,first,last in parts:
            lines,cell=excerpt(func,first,last)
            if rows: rows.append((None,''))
            rows.extend(lines)
            refs.append(f'cell {cell}, lines {lines[0][0]}–{lines[-1][0]}')
        path,w,h=render(rows,name)
        source='Project_2_Dijkstra.ipynb • '+'; '.join(refs)
        s=slide(title,section,notes='Image rendered directly from the actual notebook source. '+source+'. Original indentation, comments and counter increments are preserved. '+takeaway+'\n\nExact excerpt:\n'+'\n'.join(line for _,line in rows))
        text(s,.58,1.6,12.1,.32,source,12,'475569')
        width=12.1;height=width*h/w
        if height>4.03: height=4.03;width=height*w/h
        s.shapes.add_picture(str(path), Inches(.6+(12.1-width)/2), Inches(2.0+(4.03-height)/2),width=Inches(width),height=Inches(height))
        band(s,takeaway)
        additions.setdefault(after,[]).append(prs.slides._sldIdLst[-1])
    for entry in list(prs.slides._sldIdLst): prs.slides._sldIdLst.remove(entry)
    for number,entry in enumerate(original,1):
        prs.slides._sldIdLst.append(entry)
        for extra in additions.get(number,[]): prs.slides._sldIdLst.append(extra)
    for number,s in enumerate(prs.slides,1):
        for shape in s.shapes:
            if shape.has_text_frame and shape.left>Inches(12) and shape.top>Inches(7):
                shape.text_frame.paragraphs[0].runs[0].text=f'{number:02}'
