"""Build the Lab 2 presentation from the notebook's saved chart outputs."""
from pathlib import Path
import json, base64
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / 'presentation_assets'
nb = json.loads((ROOT / 'Project_2_Dijkstra.ipynb').read_text())
charts = []
for cell in nb['cells']:
    for output in cell.get('outputs', []):
        if 'image/png' in output.get('data', {}):
            raw = output['data']['image/png']
            path = ASSETS / f'chart_{len(charts)+1:02}.png'
            path.write_bytes(base64.b64decode(''.join(raw) if isinstance(raw, list) else raw))
            charts.append(path)
assert len(charts) == 9

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY='12283E'; BLUE='176B9B'; ORANGE='D96B27'; INK='0F172A'
GRAY='475569'; LIGHT='EEF3F7'; WHITE='FFFFFF'; BORDER='DDE4EA'; CYAN='38BDF8'

def rect(slide,x,y,w,h,fill,line=None,radius=False):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                              Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
    sh.line.fill.background() if line is None else None
    if line: sh.line.color.rgb=RGBColor.from_string(line)
    sh._element.spPr.append(OxmlElement("a:effectLst"))
    if radius: sh.adjustments[0]=0.06
    return sh

def text(slide,x,y,w,h,content,size=21,color=INK,bold=False,font='Arial'):
    sh=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(9)
    return sh

def slide(title,section='SC2001 • LAB 2',dark=False,notes=''):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else 'F7F9FB')
    text(s,.55,.28,12,.3,section,11,CYAN if dark else BLUE,True)
    text(s,.55,.83,12.2,.8,title,30,WHITE if dark else NAVY,True)
    rect(s,.55,7.04,12.2,.015,'566779' if dark else BORDER)
    text(s,.55,7.14,11,.18,'SC2001  |  Project 2: Dijkstra’s algorithm',9,'C8D5DF' if dark else GRAY)
    text(s,12.15,7.1,.6,.25,f'{len(prs.slides):02}',11,'C8D5DF' if dark else GRAY)
    s.notes_slide.notes_text_frame.text=notes
    return s

def band(s,content,dark=False):
    rect(s,.55,6.22,12.2,.58,BLUE if dark else LIGHT,radius=True)
    text(s,.77,6.34,11.8,.36,content,18,WHITE if dark else NAVY,True)

def card(s,x,tag,title,body,formula=None,dark=False,w=3.91):
    rect(s,x,1.92,w,4.04,NAVY if dark else WHITE,BORDER if not dark else None,True)
    text(s,x+.24,2.17,w-.48,.3,tag,12,CYAN if dark else BLUE,True)
    text(s,x+.24,2.68,w-.48,.67,title,23,WHITE if dark else NAVY,True)
    text(s,x+.24,3.48,w-.48,1.36,body,19,'C8D5DF' if dark else GRAY)
    if formula:
        rect(s,x+.24,5.13,w-.48,.56,BLUE if dark else LIGHT,radius=True)
        text(s,x+.36,5.24,w-.72,.38,formula,17 if len(formula)>20 else 20,WHITE if dark else BLUE,True)

def chart_slide(title,section,index,headline,body,takeaway,notes):
    s=slide(title,section,notes=notes)
    # Keep one notebook plot per slide, with the linear axes intact.
    from PIL import Image
    im=Image.open(charts[index]);iw,ih=im.size
    w=8.45;h=w*ih/iw
    if h>4.6: h=4.6;w=h*iw/ih
    s.shapes.add_picture(str(charts[index]),Inches(.42+(8.55-w)/2),Inches(1.57+(4.6-h)/2),width=Inches(w),height=Inches(h))
    rect(s,9.13,1.86,3.6,4.14,WHITE,BORDER,True)
    text(s,9.38,2.13,3.1,.79,headline,25,BLUE,True)
    text(s,9.38,3.15,3.1,2.5,body,19,GRAY)
    band(s,takeaway)
    return s

s=slide('SC2001 Lab 2',dark=True,notes='Introduce the project: the algorithm remains Dijkstra, while the graph representation and priority queue change. This deck uses the same authors shown in the supplied Lab 1 deck. All measured results are operation counts, not elapsed-time speedups.')
text(s,.6,2.02,11.8,1.1,'Dijkstra’s Algorithm',44,WHITE,True)
text(s,.63,3.22,11.6,.7,'Graph representation • Priority queues • Empirical complexity',24,'C8D5DF')
text(s,.63,4.48,10,1.25,'JAY LI YIJIE (U2540332B)\nJOEL LOH QUAN YU (U2421696K)\nJUSTIN HOTASI (U2310560D)',17,WHITE)
band(s,'Same shortest paths. Different amounts of work.',True)

s=slide('Two implementation choices affect the work',notes='The project brief prescribes matrix plus array in part A and adjacency list plus minimizing heap in part B. Neither representation is required for Dijkstra correctness. The goal is to compare their costs, not to claim one always wins.')
card(s,.55,'PART A','Matrix + array','Scan possible neighbors.\nFind the minimum by a linear queue scan.','Θ(V²)')
card(s,4.69,'PART B','List + min-heap','Visit existing neighbors.\nExtract minima and update priorities with a heap.','O((V + E) log V)')
card(s,8.83,'RESEARCH QUESTION','When is each better?','V = number of vertices\nE = directed edges\nR = successful updates','Compare counted work',True)
band(s,'Compare the same graph and source; change only the implementation.')

s=slide('Build one graph; convert it two ways','GRAPH REPRESENTATION',notes='The edge list is the canonical input. Both conversions preserve the same directed weighted edges. None means absent, so zero remains a valid edge weight. Both outer containers are Python lists: indexing vertex u takes constant time. A list of lists does not require searching for u.')
text(s,.6,1.75,12,.5,'Example: 0 → 1 (weight 4), 0 → 2 (weight 2), 1 → 2 (weight 1)',22,NAVY)
rect(s,.55,2.57,5.95,3.15,WHITE,BORDER,True);rect(s,6.77,2.57,5.95,3.15,WHITE,BORDER,True)
text(s,.82,2.8,5.4,.4,'ADJACENCY MATRIX',15,BLUE,True)
text(s,.82,3.4,5.4,1.45,'[None, 4,    2   ]\n[None, None, 1   ]\n[None, None, None]',21,NAVY,font='Courier New')
text(s,.82,5.03,5.4,.48,'Scan V cells in row u.',21,GRAY)
text(s,7.04,2.8,5.4,.4,'LIST OF ADJACENCY LISTS',15,BLUE,True)
text(s,7.04,3.4,5.4,1.45,'[[(1, 4), (2, 2)],\n [(2, 1)],\n []]',21,NAVY,font='Courier New')
text(s,7.04,5.03,5.4,.48,'Visit only u’s outgoing edges.',21,GRAY)
band(s,'matrix[u] and adjacency[u] are both O(1) accesses.')

s=slide('Both implementations follow the lecture pseudocode','ALGORITHM • SLIDES 9 AND 16',notes='Initialize parents to None and all distances to infinity, except the source at zero. Insert every vertex initially. The queue is drained completely, including unreachable vertices. With nonnegative edges, relaxation cannot improve a finalized vertex. Heap priorities must be repaired whenever a tentative distance decreases.')
rect(s,.55,1.87,7.5,4.12,NAVY,radius=True)
text(s,.85,2.13,6.9,3.5,'Initialize(G, source)\nQ ← all vertices\nwhile Q is not empty:\n    u ← ExtractMin(Q)\n    for each outgoing edge (u, v):\n        Relax(u, v)',22,WHITE,font='Courier New')
text(s,8.49,2.13,4.1,.5,'RELAXATION',15,BLUE,True)
text(s,8.49,2.81,4.1,2.52,'If dist[u] + weight < dist[v]:\n\nUpdate dist[v] and parent[v].\n\nIn B, also call decrease-key.',21,GRAY)
band(s,'Requirements: finite, nonnegative weights; zero-weight edges are allowed.')

s=slide('Measure the work that explains the complexity','EXPERIMENTAL METHOD',notes='A core-operation score counts constant-time algorithmic events. We include work inside the queues, not just relaxation calls. The score groups some constant-size actions, such as a distance/parent update, into one event. Different events have different real costs. Do not call the ratios runtime speedups. Graph generation, conversion and validation are excluded.')
card(s,.55,'COUNT','Neighbor discovery','Matrix-cell probes in A.\nAdjacency entries in B.','V² versus E')
card(s,4.69,'COUNT','Queue + relaxation','Minimum comparisons, heap comparisons/swaps, relaxation checks and updates.','Count queue work')
card(s,8.83,'INTERPRET','Work, not seconds','Fixed seeds give reproducible counts. Constants still affect actual runtime.','No runtime claim',True)
band(s,'Counting only relaxations would hide the difference: both attempt E relaxations.')

s=slide('Vary vertices and edges separately','INPUT GENERATION AND VALIDATION',notes='Three seeds: 11, 29 and 47. Directed simple graphs, no self-loops, weights 1 to 100. A shuffled chain from source zero guarantees reachability. Other edges are sampled without replacement. This is a controlled family, not a uniform sample of all graphs. The V sweeps contain 30 graphs; fixed-V E sweep adds 27. Error bars are one standard deviation across seeds. Separate tests cover zero weights and unreachable vertices.')
card(s,.55,'01 • V SWEEP','Grow the graph','V: 24, 48, 96, 192, 384\nSparse: E = 4V\nDense: E = ⌊V(V − 1)/2⌋','3 seeds per setting')
card(s,4.69,'02 • E SWEEP','Increase density','Fix V = 192.\nIncrease E from 191 to 36,672.\nKeep source = 0.','Identical inputs')
card(s,8.83,'03 • VALIDATE','Independent oracle','Compare against Bellman–Ford. Check reconstructed paths and invalid inputs.','57 experiment graphs',True)
band(s,'All presentation plots use linear axes and one graph per slide.')

s=slide('Part A: scan the queue, then scan a matrix row','PART A • IMPLEMENTATION',notes='The array queue is unsorted. Find the minimum by comparing each remaining candidate. Remove by replacing it with the last queue entry and popping the last item, which avoids a linear shift. Priorities live in dist, so updates cost constant time. Next scan all V columns of row u, even if most entries are absent.')
card(s,.55,'STEP 01','Extract minimum','Scan the remaining array entries; choose the smallest dist.','O(|Q|) per extraction')
card(s,4.69,'STEP 02','Discover neighbors','Read matrix[u][v] for every v. Relax only present edges.','V probes per vertex')
card(s,8.83,'STEP 03','Update distance','Write dist[v] and parent[v]. Queue reads the updated distance directly.','O(1) per update',True)
band(s,'Matrix = edge discovery. Array queue = choice of the next vertex.')

s=slide('Part A: two sources of quadratic work','PART A • THEORETICAL ANALYSIS',notes='Minimum selection makes V(V−1)/2 comparisons because the remaining queue sizes are V down to 1. Matrix traversal performs V² probes. Relaxation checks cost E and successful updates R≤E. Initialization is linear. Since E≤V(V−1), the result is Θ(V²). This bound also applies to disconnected inputs because every vertex is processed.')
card(s,.55,'STEP 01','Queue selection','(V − 1) + (V − 2) + … + 0 comparisons.','V(V − 1)/2')
card(s,4.69,'STEP 02','Matrix traversal','V rows, each with V cells.\nPlus E relaxation checks and R successful updates.','V² + E + R')
card(s,8.83,'CONCLUSION','Total complexity','E ≤ V(V − 1), and R ≤ E.\nMatrix storage: Θ(V²).\nAuxiliary storage: Θ(V).','Θ(V² + E) = Θ(V²)',True)
band(s,'Even a sparse graph pays for all V² possible-neighbor checks.')

chart_slide('Part A: sparse graphs still incur quadratic work','PART A • V SWEEP',0,'E = 4V','Most edges are absent.\n\nThe matrix and queue scans still grow quadratically.','Measured growth follows the V² reference.','The dashed V² reference is scaled to the first measured point. Array comparisons and matrix probes match their exact formulas on every run. Error bars represent one standard deviation over three seeds. The fitted slope, computed numerically in log space but not plotted on log axes, is 1.937.')
chart_slide('Part A: dense graphs show the same growth order','PART A • V SWEEP',1,'E ≈ V² / 2','More edges add relaxation work.\n\nThe dominant growth remains quadratic.','Dense inputs add work without changing A’s Θ(V²) bound.','Dense means E=floor(V(V−1)/2), roughly half of all possible directed edges. The measured finite-range fitted slope is 1.974. This is consistent with the exact quadratic scan counts, not an empirical proof of asymptotics.')
chart_slide('Part A: at fixed V, scan costs remain constant','PART A • E SWEEP',2,'V = 192','36,864 matrix probes.\n\n18,336 queue comparisons.\n\nMore edges add relaxations.','Increasing E adds work above an already large fixed baseline.','The exact score is 3V + V(V−1)/2 + V² + E + R. Holding V fixed leaves only E and successful updates R variable. Do not say that A is literally unaffected by E: only its simplified asymptotic bound suppresses that dependence.')

s=slide('Part B: visit existing edges and repair the heap','PART B • IMPLEMENTATION',notes='The heap contains one entry per unprocessed vertex, not duplicate lazy entries. position[v] finds its heap index in O(1). Bottom-up heap construction is O(V). Extract-min removes the root, replaces it with the last entry and sifts down. A decreased priority is restored with sift-up. Both dist arrays are shared with their queues.')
card(s,.55,'STEP 01','Extract minimum','Remove the root of an indexed binary min-heap.','O(log V)')
card(s,4.69,'STEP 02','Visit neighbors','Access adjacency[u] in O(1). Iterate its existing edges only.','Θ(outdegree(u))')
card(s,8.83,'STEP 03','Decrease key','After a successful relaxation, find v with position[v] and sift up.','O(log V) per update',True)
band(s,'The list avoids absent edges; the heap avoids linear minimum searches.')

s=slide('Part B: edge visits plus heap maintenance','PART B • THEORETICAL ANALYSIS',notes='The detailed bound is O(V+E+(V+R)log V), where R≤E. This yields O((V+E)log V). For reachable graphs E≥V−1, the slides simplify it to O(E log V). With sparse graphs E=Θ(V), the bound is O(V log V). With dense graphs the worst-case upper bound is O(V² log V), but actual updates may be far fewer than E.')
card(s,.55,'STEP 01','Visit existing edges','Each adjacency entry is read once.\nEach edge is relaxed once.','Θ(E)')
card(s,4.69,'STEP 02','Maintain the heap','V extract-min calls.\nR decrease-key calls.\nHeap construction is O(V).','O((V + R) log V)')
card(s,8.83,'CONCLUSION','Worst-case bound','Use R ≤ E.\nGraph space: Θ(V + E).\nAuxiliary space: Θ(V).','O((V + E) log V)',True)
band(s,'An upper bound does not say every edge causes a logarithmic-cost update.')

chart_slide('Part B: sparse graphs benefit from both choices','PART B • V SWEEP',3,'O(V log V)','Only 4V edges are visited.\n\nThe heap replaces repeated linear minimum searches.','Sparse inputs avoid both quadratic costs found in Part A.','The dashed line is the scaled upper-bound growth (V+E) log₂ V, normalized to the first measured point. It is a reference shape, not a guaranteed numerical upper bound. A power-law fit of 1.150 over this finite range is compatible with V log V growth.')
chart_slide('Part B: a dense upper bound can be pessimistic','PART B • V SWEEP',4,'R can be ≪ E','Every edge is checked.\n\nOnly successful improvements trigger decrease-key.','Observed work can grow closer to V² than V² log V.','Part B has a fitted finite-range slope of 1.814 here. This is not a claim of subquadratic asymptotic complexity on dense graphs: visiting E=Θ(V²) edges already gives a quadratic lower bound. The finite input range and overhead explain finite-range exponents. The reference curve is not a prediction of an exact count.')
chart_slide('Part B: adding edges directly adds list visits','PART B • E SWEEP',5,'V = 192','E list-entry visits.\n\nE relaxation checks.\n\nHeap work depends on successful updates.','At fixed V, edge traversal grows linearly with E.','The exact score is 3V + 2E + 2R + heap comparisons + heap swaps. V is fixed, but R and the distance moved by keys depend on structure and weights. Increasing density does not imply every new edge successfully updates a distance.')

chart_slide('Part C: Part B uses much less work on sparse inputs','PART C • COMPARISON',6,'A / B = 15.21','At V = 384, E = 1,536:\n\nA: 224,192 operations\nB: 14,737 operations','This is a work-score ratio, not a runtime speedup.','Values are rounded means across three seeds. Blue is A and orange is B. Solid curves are measured, dashed curves are theoretical growth references: V²+E and (V+E)log₂V. Each is normalized to its own first measured point. Both graph representation and queue contribute to B’s advantage.')
chart_slide('Part C: Part B also wins these random dense inputs','PART C • COMPARISON',7,'A / B = 1.82','At V = 384, E = 73,536:\n\nA: 297,284 operations\nB: 163,496 operations','A’s better worst-case bound does not guarantee a win on every dense graph.','Retain this result rather than claiming density automatically makes A faster. Dense random graphs can have far fewer successful updates than edges. The theoretical lines compare growth after independent first-point scaling; their crossing is not a universal crossover prediction.')
chart_slide('Part C: why the ratio falls as E increases','PART C • FIXED V',8,'Read A / B','Above 1: B uses less work.\n\nBelow 1: A uses less work.\n\nCrossing 1 is not required.','A’s fixed scan costs are large; B’s edge-dependent work catches up.','At V=192, A always performs V² matrix checks and V(V−1)/2 queue comparisons. B initially visits few edges, then more as E grows. Therefore B’s advantage shrinks on these datasets. Both algorithms also perform E relaxations. The dashed curve is a ratio of independently scaled theoretical references, not a proven bound on the measured ratio. All axes are linear.')

s=slide('Frequent updates can make Part A preferable','PART C • A CONTROLLED DENSE EXAMPLE',notes='This is a separate constructed workload, not part of the random trials. It is a complete directed graph at V=384 and E=147072. Chain edges u→u+1 have weight 1; other forward edges have weight 2V−2u+v; backward edges have weight 2V. Then dist[u]=u and every forward edge improves its target, so R=E/2=73536. Both implementations are checked against Bellman–Ford. A has about 15.3% lower score than B. The example does not prove every decrease-key costs Θ(log V).')
text(s,.65,1.75,11.8,.6,'Complete directed graph • V = 384 • E = 147,072 • R = E / 2',21,GRAY)
for y,name,value,color in [(2.75,'Part A',442752,BLUE),(4.1,'Part B',522916,ORANGE)]:
    text(s,.67,y+.08,1.3,.45,name,24,NAVY,True)
    rect(s,2.1,y,7.5*value/550000,.69,color)
    text(s,2.1+7.5*value/550000+.16,y+.1,2.3,.4,f'{value:,}',23,color,True)
text(s,2.1,5.23,9,.45,'Counted operations; bars start at zero.',17,GRAY)
band(s,'Part A uses 15.3% fewer counted operations when many priorities improve.')

s=slide('Separate the representation from the priority queue','PART C • WHAT ACTUALLY CAUSES THE DIFFERENCE?',notes='This slide addresses a common misconception. A list can be paired with an array queue and retain Θ(V²+E) time, while visiting only existing edges. Thus the matrix itself is not the source of A’s better dense-graph worst-case guarantee. A matrix offers O(1) lookup of a particular edge; a list of lists needs a neighbor search. Dijkstra usually needs neighbor iteration instead. Representation and queue are changed together in the assignment, so the experiment cannot isolate each choice causally.')
rows=[('Choice','Matrix / array','List / heap'),('Neighbor discovery','Matrix: V² cells','List: E entries'),('Minimum selection','Array: Θ(V²) total','Heap: O(V log V) total'),('Priority updates','Array: O(1) each','Heap: O(log V) each'),('Specific edge lookup','Matrix: O(1)','List: O(outdegree(u))')]
for i,row in enumerate(rows):
    y=1.92+i*.76
    rect(s,.55,y,12.2,.7,NAVY if i==0 else WHITE)
    for j,(x,w) in enumerate([(.77,3.4),(4.4,3.8),(8.43,4.0)]):
        text(s,x,y+.16,w,.41,row[j],19,WHITE if i==0 else NAVY,i==0)
band(s,'List + array also gives Θ(V² + E): a matrix is not essential for that bound.')

s=slide('Choose based on sparsity and update behavior','CONCLUSIONS',dark=True,notes='Close with three messages. First, B is the natural sparse-graph choice: both traversal and selection avoid quadratic work. Second, A has a stronger dense-graph worst-case bound, but measured results depend on updates and constants. Third, operation counts explain growth but not elapsed runtime. The assignment fixes the combinations; do not silently replace its matrix with a list. Invite questions.')
text(s,.7,1.95,1,.5,'01',25,CYAN,True);text(s,1.65,1.95,10.7,1.0,'Sparse graphs: list + heap avoids wasted scans and saves space.',27,WHITE,True)
text(s,.7,3.35,1,.5,'02',25,CYAN,True);text(s,1.65,3.35,10.7,1.0,'Dense graphs: array updates give a strong worst-case bound; a win is not automatic.',27,WHITE,True)
text(s,.7,4.92,1,.5,'03',25,CYAN,True);text(s,1.65,4.92,10.7,1.0,'Report what was measured: core-operation counts, not CPU speedups.',27,WHITE,True)
band(s,'Thank you — questions?',True)

s=slide('Exact counters connect theory to the measurements','APPENDIX • COUNTING MODEL',notes='Count initializations, queue entries, extractions, array comparisons, matrix probes, list visits, relaxation checks, distance/parent updates, heap comparisons, heap swaps and decrease-key calls. Each is one declared event; these are not equally costly CPU instructions. Hc and Hs include heap construction. Assertions verify both exact score identities on every experiment.')
text(s,.72,1.9,11.9,.5,'Part A',23,BLUE,True)
text(s,.72,2.57,11.9,.6,'C_A = 3V + V(V − 1)/2 + V² + E + R',29,NAVY)
text(s,.72,3.55,11.9,.5,'Part B',23,ORANGE,True)
text(s,.72,4.22,11.9,.6,'C_B = 3V + 2E + 2R + H_c + H_s',29,NAVY)
text(s,.72,5.22,11.8,.65,'R: successful relaxations     H_c: heap comparisons     H_s: heap swaps',19,GRAY)
band(s,'Generation, conversion, plotting and validation are outside these counts.')

s=slide('Validation, scope and sources','APPENDIX • REPRODUCIBILITY',notes='The notebook has six hand-crafted correctness cases, 40 random correctness graphs, 57 experimental graphs and six invalid input cases. Four additional high-update dense graphs are checked as well. References: Project 2.pdf parts a–c; 06_ShortestPath.pdf printed slides 9–10 for relaxation and initialization, 13 for Bellman–Ford, 15–18 for Dijkstra. The slide deck is based on the saved notebook outputs. The original Lab 1 deck supplies only design and author information, not results.')
card(s,.55,'VALIDATION','Independent checks','Bellman–Ford distances.\nReconstructed path weights.\nZero weights, ties and disconnected graphs.','All checks passed')
card(s,4.69,'LIMITATIONS','Controlled experiments','Three seeds per setting.\nModerate input sizes.\nCounts ≠ elapsed time.\nReferences are scaled.','No universal crossover')
card(s,8.83,'SOURCES','Project materials','Project 2.pdf: parts a–c\n06_ShortestPath.pdf: slides 9–10, 13, 15–18\nProject_2_Dijkstra.ipynb','Notebook results',True)
band(s,'Speaker notes include the detailed explanations and assumptions.')

from code_images import add_code_slides
add_code_slides(prs, nb, ASSETS, slide, text, band)

output=ROOT/'SC2001 Lab 2 Slides.pptx'
prs.save(output)
for n,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        assert sh.left >= 0 and sh.top >= 0, (n,'negative position')
        assert sh.left+sh.width <= prs.slide_width+10000, (n,'right overflow')
        assert sh.top+sh.height <= prs.slide_height+10000, (n,'bottom overflow')
print(f'Created {output.name}: {len(prs.slides)} slides, {len(charts)} notebook charts, speaker notes on every slide.')
